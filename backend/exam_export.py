"""
SSTG — Date-aware exam exports (PDF + Excel)
=============================================
Replaces ExcelExporter.exam_pdf / exam_pdf_layout / exam_xlsx / exam_xlsx_layout
(which grouped slots by weekday NAME and therefore merged week 1 and week 2 of
a multi-week session). One implementation now serves both the "no layout" and
"layout template" cases; exporter.py is left untouched and still provides the
shared header/theme helpers.
"""
from __future__ import annotations

from collections import defaultdict
from datetime import date
from io import BytesIO

from fastapi import HTTPException

from exporter import ExcelExporter

DAY_ORDER = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]


class _DefaultLayout:
    name = None
    group_by = "day"
    orientation = "landscape"
    show_duration = True
    show_invigilator = True
    show_room = True
    show_notes = False
    show_practical_tag = True
    footer_text = None
    warning_text = None


def _pretty_date(iso: str | None, day: str) -> str:
    if not iso:
        return day
    try:
        return date.fromisoformat(iso).strftime("%a %d %b %Y")
    except ValueError:
        return f"{day} {iso}"


def _sort_key(sl):
    di = DAY_ORDER.index(sl.day) if sl.day in DAY_ORDER else 9
    return (sl.exam_date or "9999-99-99", di, sl.period, sl.class_section.name)


class ExamExporter(ExcelExporter):

    def _layout_for(self, db, layout):
        """No explicit layout → mirror the school-wide export column toggles."""
        if layout is not None:
            return layout
        inc_sup, inc_room = self._exam_cols(db)
        lay = _DefaultLayout()
        lay.show_invigilator, lay.show_room = bool(inc_sup), bool(inc_room)
        return lay

    def _slots(self, session, db):
        from models import ExamSlot
        rows = db.query(ExamSlot).filter(ExamSlot.exam_session_id == session.id).all()
        return sorted(rows, key=_sort_key)

    # ── PDF ───────────────────────────────────────────────────────────────────

    def exam_pdf_unified(self, session, db, layout=None) -> bytes:
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.units import cm
            from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
        except ImportError:
            raise HTTPException(503, "reportlab not installed. Run: pip install reportlab")

        lay = self._layout_for(db, layout)
        page = landscape(A4) if lay.orientation == "landscape" else A4
        buf = BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=page, leftMargin=1*cm, rightMargin=1*cm,
                                topMargin=1.5*cm, bottomMargin=1.5*cm)
        styles = getSampleStyleSheet()
        story = []
        suffix = f"  ({lay.name})" if lay.name else ""
        self._header(story, f"Exam Timetable — {session.name}{suffix}", db)
        story.append(Paragraph(f"{session.start_date} to {session.end_date}",
                               ParagraphStyle("S", parent=styles["Normal"], fontSize=9,
                                              textColor=colors.grey, spaceAfter=8)))
        if lay.warning_text:
            story.append(Paragraph(f"\u26a0 {lay.warning_text}", ParagraphStyle(
                "W", parent=styles["Normal"], fontSize=9.5, textColor=colors.HexColor("#92400e"),
                backColor=colors.HexColor("#fef3c7"), borderPadding=6, spaceAfter=10,
                fontName="Helvetica-Bold")))

        slots = self._slots(session, db)
        if not slots:
            story.append(Paragraph("No exam slots scheduled yet.", styles["Normal"]))
            doc.build(story)
            return buf.getvalue()

        by_class = lay.group_by == "class"
        groups = defaultdict(list)
        order = []
        for sl in slots:
            key = sl.class_section.name if by_class else _pretty_date(sl.exam_date, sl.day)
            if key not in groups:
                order.append(key)
            groups[key].append(sl)
        if by_class:
            order = sorted(order)

        head_style = ParagraphStyle("D", parent=styles["Heading2"],
                                    textColor=colors.HexColor("#1E3A5F"), spaceAfter=4, spaceBefore=10)
        total_w = (24 if lay.orientation == "landscape" else 17) * cm

        for gi, key in enumerate(order):
            story.append(Paragraph(key, head_style))
            header = ["Date · Period" if by_class else "Period", "Subject", "Paper"]
            if lay.show_duration:
                header.append("Duration")
            if not by_class:
                header.append("Class")
            if lay.show_invigilator:
                header.append("Invigilator")
            if lay.show_room:
                header.append("Room")
            if lay.show_notes:
                header.append("Notes")
            rows = [header]
            for sl in groups[key]:
                first = (f"{_pretty_date(sl.exam_date, sl.day)} · P{sl.period}" if by_class else f"P{sl.period}")
                paper = f"Paper {sl.paper.paper_number}" + ("*" if lay.show_practical_tag and sl.paper.is_practical else "")
                row = [first, sl.paper.subject.name, paper]
                if lay.show_duration:
                    row.append(f"{sl.paper.duration_minutes} min")
                if not by_class:
                    row.append(sl.class_section.name)
                if lay.show_invigilator:
                    row.append(sl.invigilator.name if sl.invigilator else "\u2014")
                if lay.show_room:
                    row.append(sl.room or "\u2014")
                if lay.show_notes:
                    row.append(sl.notes or "")
                rows.append(row)
            tbl = Table(rows, colWidths=[total_w / len(header)] * len(header), repeatRows=1)
            tbl.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A5F")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EFF6FF")]),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]))
            story.append(tbl)
            story.append(Spacer(1, 0.4*cm))
            if gi < len(order) - 1:
                story.append(PageBreak())

        if lay.show_practical_tag:
            story.append(Paragraph("* Practical paper", styles["Normal"]))
        if lay.footer_text:
            story.append(Spacer(1, 0.5*cm))
            story.append(Paragraph(lay.footer_text, ParagraphStyle(
                "F", parent=styles["Normal"], fontSize=8, textColor=colors.grey, alignment=1)))
        doc.build(story)
        return buf.getvalue()

    # ── Excel ─────────────────────────────────────────────────────────────────

    def exam_xlsx_unified(self, session, db, layout=None) -> bytes:
        from exporter import XLSX
        if not XLSX:
            raise HTTPException(503, "openpyxl not installed. Run: pip install openpyxl")
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        from exporter import _thin_border, _EXAM_H, _WHITE, _GREY_L

        lay = self._layout_for(db, layout)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = ((lay.name or "Exam Schedule")[:28]) or "Exam Schedule"

        headers = ["Date", "Day", "Period", "Class", "Subject", "Paper"]
        if lay.show_duration:
            headers.append("Duration (min)")
        if lay.show_invigilator:
            headers.append("Invigilator")
        if lay.show_room:
            headers.append("Room")
        if lay.show_notes:
            headers.append("Notes")
        last = get_column_letter(len(headers))

        ws["A1"] = self._school_info(db)["name"]
        ws["A1"].font = Font(bold=True, size=16, color=_EXAM_H)
        ws["A2"] = f"Exam Session: {session.name}   |   {session.start_date} \u2192 {session.end_date}"
        ws["A2"].font = Font(italic=True, size=11, color="555555")
        ws.merge_cells(f"A1:{last}1")
        ws.merge_cells(f"A2:{last}2")

        row = 4
        if lay.warning_text:
            c = ws.cell(row=row, column=1, value=f"\u26a0 {lay.warning_text}")
            c.font = Font(bold=True, size=10, color="92400E")
            c.fill = PatternFill("solid", fgColor="FEF3C7")
            row += 2
        for ci, h in enumerate(headers, 1):
            c = ws.cell(row=row, column=ci, value=h)
            c.font = Font(bold=True, color=_WHITE)
            c.fill = PatternFill("solid", fgColor=_EXAM_H)
            c.alignment = Alignment(horizontal="center")
            c.border = _thin_border()
        row += 1

        for n, sl in enumerate(self._slots(session, db)):
            paper = f"Paper {sl.paper.paper_number}" + ("*" if lay.show_practical_tag and sl.paper.is_practical else "")
            vals = [sl.exam_date or "", sl.day, sl.period, sl.class_section.name, sl.paper.subject.name, paper]
            if lay.show_duration:
                vals.append(sl.paper.duration_minutes)
            if lay.show_invigilator:
                vals.append(sl.invigilator.name if sl.invigilator else "")
            if lay.show_room:
                vals.append(sl.room or "")
            if lay.show_notes:
                vals.append(sl.notes or "")
            for ci, v in enumerate(vals, 1):
                c = ws.cell(row=row, column=ci, value=v)
                c.alignment = Alignment(horizontal="center")
                c.border = _thin_border()
                if n % 2:
                    c.fill = PatternFill("solid", fgColor=_GREY_L)
            row += 1

        if lay.footer_text:
            ws.cell(row=row + 1, column=1, value=lay.footer_text).font = Font(italic=True, size=9, color="777777")
        for ci, w in enumerate([13, 12, 8, 10, 22, 12, 14, 22, 16, 20][:len(headers)], 1):
            ws.column_dimensions[get_column_letter(ci)].width = w

        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()
